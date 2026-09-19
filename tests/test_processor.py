"""测试处理器对 DOCX、XLSX、TXT 的脱敏与输出
"""

import os
import unittest
import tempfile
import docx
import openpyxl
from gearify.core.desensitizer import Desensitizer, DesensitizerConfig
from gearify.core.processor import FileProcessor, get_output_path


class TestFileProcessor(unittest.TestCase):

    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp()
        config = DesensitizerConfig(
            custom_names=["李明", "诸葛亮"],
            custom_units=["阿里巴巴", "腾讯公司"],
            custom_industries=["人工智能", "金融"],
        )
        self.desensitizer = Desensitizer(config)
        self.processor = FileProcessor(self.desensitizer)

    def test_text_processing(self):
        txt_path = os.path.join(self.tmp_dir, "sample.txt")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write("一、总则\n1. 项目负责人是李明，总投资1500万元，属于人工智能行业。\n")

        out_path, stats = self.processor.process_single_file(txt_path)
        self.assertTrue(os.path.exists(out_path))
        self.assertTrue(out_path.endswith("sample_脱敏.txt"))

        with open(out_path, "r", encoding="utf-8-sig") as f:
            content = f.read()

        self.assertIn("一、总则", content)
        self.assertIn("1. 项目负责人是××", content)
        self.assertIn("总投资××万元", content)
        self.assertIn("属于×××行业", content)

    def test_docx_processing(self):
        doc_path = os.path.join(self.tmp_dir, "sample.docx")
        doc = docx.Document()
        doc.add_heading("第一章 项目立项", level=1)
        doc.add_paragraph("联系人：诸葛亮，身份证：420102199001011234，来自腾讯公司。")
        
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "部门"
        table.cell(0, 1).text = "预算"
        table.cell(1, 0).text = "阿里巴巴"
        table.cell(1, 1).text = "5000万元"
        doc.save(doc_path)

        out_path, stats = self.processor.process_single_file(doc_path)
        self.assertTrue(os.path.exists(out_path))
        self.assertTrue(out_path.endswith("sample_脱敏.docx"))

        res_doc = docx.Document(out_path)
        p1 = res_doc.paragraphs[0].text
        p2 = res_doc.paragraphs[1].text
        self.assertEqual(p1, "第一章 项目立项")
        self.assertIn("诸葛亮", "联系人：诸葛亮") # Original
        self.assertIn("联系人：×××", p2)
        self.assertIn("身份证：×××", p2)
        self.assertIn("来自×××。", p2)

        tbl = res_doc.tables[0]
        self.assertEqual(tbl.cell(1, 0).text, "×××")
        self.assertIn("××万元", tbl.cell(1, 1).text)

    def test_xlsx_processing(self):
        xlsx_path = os.path.join(self.tmp_dir, "sample.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "姓名"
        ws["B1"] = "金额"
        ws["A2"] = "李明"
        ws["B2"] = 1200000
        wb.save(xlsx_path)

        out_path, stats = self.processor.process_single_file(xlsx_path)
        self.assertTrue(os.path.exists(out_path))
        self.assertTrue(out_path.endswith("sample_脱敏.xlsx"))

        res_wb = openpyxl.load_workbook(out_path)
        res_ws = res_wb.active
        self.assertEqual(res_ws["A2"].value, "××")
        self.assertEqual(res_ws["B2"].value, "××")

    def test_docx_run_style_preservation(self):
        doc_path = os.path.join(self.tmp_dir, "styled.docx")
        doc = docx.Document()
        p = doc.add_paragraph()
        r1 = p.add_run("联系人：")
        r2 = p.add_run("李明")
        r2.bold = True
        r3 = p.add_run("，电话：")
        r4 = p.add_run("13812345678")
        r4.italic = True
        doc.save(doc_path)

        out_path, _ = self.processor.process_single_file(doc_path)
        res_doc = docx.Document(out_path)
        res_p = res_doc.paragraphs[0]
        self.assertEqual(len(res_p.runs), 4)
        self.assertEqual(res_p.runs[0].text, "联系人：")
        self.assertEqual(res_p.runs[1].text, "××")
        self.assertTrue(res_p.runs[1].bold)
        self.assertEqual(res_p.runs[2].text, "，电话：")
        self.assertEqual(res_p.runs[3].text, "×××")
        self.assertTrue(res_p.runs[3].italic)

    def test_xlsx_formula_and_sequence_protection(self):
        xlsx_path = os.path.join(self.tmp_dir, "advanced.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["序号", "手机号", "预算", "合计公式"])
        ws.append([1, 13812345678, 5000, "=SUM(C2:C10)"])
        ws.append([2, 13987654321, 6000, "=100+200"])
        wb.save(xlsx_path)

        out_path, _ = self.processor.process_single_file(xlsx_path)
        res_wb = openpyxl.load_workbook(out_path)
        res_ws = res_wb.active

        # 序号列保护
        self.assertEqual(res_ws["A2"].value, 1)
        self.assertEqual(res_ws["A3"].value, 2)

        # 手机长整数脱敏为实体 ×××
        self.assertEqual(res_ws["B2"].value, "×××")
        self.assertEqual(res_ws["B3"].value, "×××")

        # 预算数字脱敏
        self.assertEqual(res_ws["C2"].value, "××")
        self.assertEqual(res_ws["C3"].value, "××")

        # 公式保护：未被替换为字符或被脱敏
        self.assertEqual(res_ws["D2"].value, "=SUM(C2:C10)")
        self.assertEqual(res_ws["D3"].value, "=100+200")


if __name__ == "__main__":
    unittest.main()
